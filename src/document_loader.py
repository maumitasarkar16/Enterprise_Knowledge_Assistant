"""Local TXT, Markdown, PDF and DOCX document ingestion."""

import logging
from pathlib import Path
from pypdf import PdfReader
from docx import Document as DocxDocument
from src.models import Document

logger = logging.getLogger(__name__)
SUPPORTED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}

def _load_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    pages = []
    for page_no, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[Page {page_no}]\n{text}")
    return "\n\n".join(pages)

def _load_docx(path: Path) -> str:
    doc = DocxDocument(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

def load_documents(data_dir: str) -> list[Document]:
    root = Path(data_dir)
    if not root.exists():
        raise FileNotFoundError(f"Document directory does not exist: {root.resolve()}")

    documents = []
    for path in sorted(root.rglob("*")):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            ext = path.suffix.lower()
            if ext in {".txt", ".md"}:
                text = path.read_text(encoding="utf-8")
            elif ext == ".pdf":
                text = _load_pdf(path)
            else:
                text = _load_docx(path)

            if not text.strip():
                logger.warning("Skipping empty document: %s", path)
                continue

            documents.append(
                Document(
                    text=text.strip(),
                    metadata={
                        "source": path.name,
                        "path": str(path),
                        "extension": ext,
                    },
                )
            )
        except Exception:
            logger.exception("Failed to load %s; continuing", path)

    if not documents:
        raise ValueError(f"No readable supported documents found under {root.resolve()}")
    return documents
